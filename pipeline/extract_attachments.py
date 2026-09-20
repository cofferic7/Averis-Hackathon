"""
extract_attachments.py
Turn any SI / BL attachment (txt, xlsx, docx, pdf) into plain text.

Every call to extract() returns the SAME shape, and never raises:
    {"path", "text", "method", "status": "ok" | "needs_review", "warning"}

Install:
    pip install openpyxl python-docx pdfplumber google-genai
Set the key for the Gemini fallback (scanned PDFs):
    export GEMINI_API_KEY="..."

Usage:
    python extract_attachments.py path/to/attachments          # a folder
    python extract_attachments.py a.pdf b.docx                 # or files
Results are cached as JSON in ./extracted/ so Gemini is not called twice.
"""
from __future__ import annotations

# pylint: disable=missing-function-docstring

import json
import os
import sys
from dataclasses import asdict, dataclass
from pathlib import Path
import openpyxl
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from google import genai
from google.genai import types
import pdfplumber

MIN_PDF_CHARS = 50            # below this, treat the PDF as scanned / image-only
GEMINI_MODEL = "gemini-2.5-flash"   # change if your account uses another model
OUT_DIR = Path("extracted")


@dataclass
class ExtractResult:
    path: str
    text: str = ""
    method: str = ""
    status: str = "ok"            # ok | needs_review
    warning: str | None = None


# --------------------------------------------------------------------------
# txt
# --------------------------------------------------------------------------
def extract_txt(path: str) -> str:
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "utf-8", "latin-1"):   # latin-1 never fails
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return ""


# --------------------------------------------------------------------------
# xlsx
# --------------------------------------------------------------------------
def extract_xlsx(path: str) -> str:

    wb = openpyxl.load_workbook(path, data_only=True)   # data_only: values, not formulas
    lines: list[str] = []
    for ws in wb.worksheets:
        lines.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if any(c is not None and str(c).strip() for c in row):
                lines.append(" | ".join("" if c is None else str(c).strip() for c in row))
    return "\n".join(lines)


# --------------------------------------------------------------------------
# docx  (paragraphs AND tables, kept in document order)
# --------------------------------------------------------------------------
def extract_docx(path: str) -> str:

    doc = Document(path)
    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = Paragraph(child, doc).text.strip()
            if text:
                lines.append(text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, doc).rows:
                cells: list[str] = []
                for cell in row.cells:
                    t = cell.text.strip().replace("\n", " ")
                    if not cells or t != cells[-1]:   # merged cells repeat text; drop repeats
                        cells.append(t)
                if any(cells):
                    lines.append(" | ".join(cells))
    return "\n".join(lines)


# --------------------------------------------------------------------------
# pdf  (text layer first, Gemini only when the text layer is empty)
# --------------------------------------------------------------------------
def extract_pdf_text_layer(path: str) -> str:

    with pdfplumber.open(path) as pdf:
        return "\n".join((p.extract_text() or "") for p in pdf.pages)


GEMINI_PROMPT = """Transcribe ALL text in this document exactly as written.
Rules:
- Preserve label/value pairs and keep table rows on one line, cells separated by " | ".
- Do not summarise, translate, correct, or infer anything.
- If a word or number is unreadable, write [ILLEGIBLE] instead of guessing.
- Output only the transcription."""


def extract_pdf_with_gemini(path: str) -> str:

    client = genai.Client(api_key=os.environ["GEMINI_API_KEY"])
    resp = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=[
            types.Part.from_bytes(data=Path(path).read_bytes(), mime_type="application/pdf"),
            GEMINI_PROMPT,
        ],
    )
    return resp.text or ""


def extract_pdf(path: str) -> ExtractResult:
    res = ExtractResult(path=path)
    try:
        text = extract_pdf_text_layer(path)
    except Exception as e:                       # corrupted / encrypted PDF
        text = ""
        res.warning = f"text layer failed: {e}"

    if len(text.strip()) >= MIN_PDF_CHARS:
        res.text, res.method = text, "pdfplumber"
        return res

    # empty text layer -> probably scanned -> vision LLM
    try:
        res.text = extract_pdf_with_gemini(path)
        res.method = "gemini"
        res.warning = ((res.warning or "") + " scanned PDF read by LLM; verify key fields").strip()
        if len(res.text.strip()) < MIN_PDF_CHARS:
            res.status, res.warning = "needs_review", "Gemini returned almost no text"
        elif "[ILLEGIBLE]" in res.text:
            res.status, res.warning = "needs_review", "document contains illegible values"
    except Exception as e:
        res.method, res.status = "gemini", "needs_review"
        res.warning = f"PDF unreadable, Gemini fallback failed: {e}"
    return res


# --------------------------------------------------------------------------
# dispatcher
# --------------------------------------------------------------------------
def extract(path: str | None) -> ExtractResult:
    if not path or not Path(path).exists():
        return ExtractResult(path=str(path), status="needs_review",
                             method="none", warning="attachment missing")

    ext = Path(path).suffix.lower()
    try:
        if ext == ".pdf":
            return extract_pdf(path)
        handlers = {".txt": extract_txt, ".xlsx": extract_xlsx, ".docx": extract_docx}
        if ext not in handlers:
            return ExtractResult(path=path, status="needs_review", method="none",
                                 warning=f"unsupported file type: {ext}")
        text = handlers[ext](path)
        res = ExtractResult(path=path, text=text, method=ext.lstrip("."))
        if not text.strip():
            res.status, res.warning = "needs_review", "file parsed but contained no text"
        return res
    except Exception as e:                       # never crash the pipeline
        return ExtractResult(path=path, status="needs_review",
                             method=ext.lstrip("."), warning=f"extraction error: {e}")


# --------------------------------------------------------------------------
# CLI: extract a folder or list of files, cache results as JSON
# --------------------------------------------------------------------------
def main(args: list[str]) -> None:
    files: list[Path] = []
    for a in args:
        p = Path(a)
        files += sorted(f for f in p.rglob("*") if f.is_file()) if p.is_dir() else [p]

    OUT_DIR.mkdir(exist_ok=True)
    for f in files:
        cache = OUT_DIR / (f.name + ".json")
        if cache.exists():                       # don't pay for Gemini twice
            continue
        result = asdict(extract(str(f)))
        cache.write_text(json.dumps(result, indent=2, ensure_ascii=False))
        print(f"{result['status']:<13} {result['method']:<10} {f.name}"
              + (f"  <- {result['warning']}" if result["warning"] else ""))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1:])